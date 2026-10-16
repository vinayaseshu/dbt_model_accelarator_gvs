#!/usr/bin/env python3
"""
dbt Model Accelerator  (v2 - best-practice edition)
===================================================
Generates a dbt project slice from a Source-to-Target Mapping (STM) spreadsheet:

  models/staging/<source>/_<source>__sources.yml        source definitions (+ optional freshness)
  models/staging/<source>/stg_<source>__<table>.sql     OPTIONAL staging layer, one per source table
  models/staging/<source>/_<source>__models.yml         OPTIONAL staging docs
  models/<folder>/<model>.sql                           one per MODEL_NAME (default folder: marts)
  models/<folder>/_<folder>__models.yml                 docs, data_type, tests, contracts
  packages.yml, dbt_project_snippet.yml, .sqlfluff      project scaffolding
  generation_report.md                                  validation + lineage report

Mapping columns (case-insensitive; "SORUCE_TABLE" typo accepted)
  Required : SOURCE_TABLE, SOURCE_COLUMN, TARGET_COLUMN, MODEL_NAME
  Core     : TARGET_TABLE, TARGET_DATATYPE, MATERIALIZATION, UNIQUE_KEY, JOIN_CONDITION
  Optional : JOIN_TYPE, INCREMENTAL_FILTER, SOURCE_NAME, MODEL_FOLDER, SCHEMA, TAGS,
             MODEL_DESCRIPTION, COLUMN_DESCRIPTION, COLUMN_TESTS, USE_STAGING

Staging layer is OPTIONAL and off by default:
  --staging            build stg_ models for every target model
  USE_STAGING = Y / N  per-model override in the mapping (blank = CLI default)

Usage
  python generate_dbt_models.py mapping.xlsx -o dbt_output
  python generate_dbt_models.py mapping.xlsx -o dbt_output --staging
  python generate_dbt_models.py mapping.xlsx -o dbt_output --enforce-contracts --surrogate-key --audit-columns
  python generate_dbt_models.py mapping.xlsx -o dbt_output --check      # CI drift check
"""
from __future__ import annotations

import argparse
import csv
import filecmp
import logging
import re
import shutil
import sys
import tempfile
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path

__version__ = "2.1.0"
log = logging.getLogger("dbt_accelerator")

# --------------------------------------------------------------------------- #
# Constants
# --------------------------------------------------------------------------- #
DEFAULT_SHEET = "STM_MAPPING"
HEADER_ALIASES = {
    "SORUCE_TABLE": "SOURCE_TABLE", "SRC_TABLE": "SOURCE_TABLE", "SRC_COLUMN": "SOURCE_COLUMN",
    "TGT_TABLE": "TARGET_TABLE", "TGT_COLUMN": "TARGET_COLUMN", "TGT_DATATYPE": "TARGET_DATATYPE",
    "DATATYPE": "TARGET_DATATYPE", "MATERIALISATION": "MATERIALIZATION", "TESTS": "COLUMN_TESTS",
    "TARGET_SCHEMA": "SCHEMA", "SOURCE_SYSTEM": "SOURCE_NAME", "FOLDER": "MODEL_FOLDER",
    "STAGING": "USE_STAGING", "STAGING_LAYER": "USE_STAGING",
}
REQUIRED = ["SOURCE_TABLE", "SOURCE_COLUMN", "TARGET_COLUMN", "MODEL_NAME"]
MODEL_LEVEL = ["MATERIALIZATION", "UNIQUE_KEY", "JOIN_CONDITION", "JOIN_TYPE", "INCREMENTAL_FILTER",
               "MODEL_FOLDER", "SCHEMA", "TAGS", "MODEL_DESCRIPTION", "USE_STAGING"]
YES = {"y", "yes", "true", "1", "on"}
NO = {"n", "no", "false", "0", "off"}
VALID_MATERIALIZATIONS = {"table", "view", "incremental", "ephemeral"}
VALID_JOIN_TYPES = {"inner", "left", "right", "full", "left outer", "right outer", "full outer"}
IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_$]*$")
SNAKE_RE = re.compile(r"^[a-z_][a-z0-9_]*$")
QUALIFIED_REF_RE = re.compile(r"\b([A-Za-z_][A-Za-z0-9_]*)\s*\.\s*([A-Za-z_][A-Za-z0-9_]*)\b")


def norm_header(h) -> str:
    key = re.sub(r"\s+", "_", str(h or "").strip().upper())
    return HEADER_ALIASES.get(key, key)


def clean(v) -> str:
    return "" if v is None else str(v).strip()


def split_list(text: str) -> list[str]:
    """Split on commas that are not inside parentheses."""
    out, depth, cur = [], 0, ""
    for ch in text:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == "," and depth == 0:
            out.append(cur.strip())
            cur = ""
        else:
            cur += ch
    if cur.strip():
        out.append(cur.strip())
    return [x for x in out if x]


# --------------------------------------------------------------------------- #
# Options
# --------------------------------------------------------------------------- #
@dataclass
class Options:
    source_name: str = "raw"
    source_schema: str | None = None
    staging: bool = False  # default for models whose USE_STAGING is blank
    default_folder: str = "marts"
    enforce_contracts: bool = False
    surrogate_key: bool = False
    audit_columns: bool = False
    composite_test: bool = True
    legacy_test_syntax: bool = False
    freshness_field: str | None = None
    project_name: str = "my_dbt_project"
    dialect: str = "snowflake"
    mapping_file: str = "mapping"

    @property
    def tests_key(self) -> str:
        return "tests" if self.legacy_test_syntax else "data_tests"


# --------------------------------------------------------------------------- #
# Data model
# --------------------------------------------------------------------------- #
@dataclass
class ColumnMap:
    row: int
    source_name: str
    source_table: str
    source_column: str
    target_table: str
    target_column: str
    target_datatype: str
    description: str = ""
    tests: str = ""

    @property
    def alias(self) -> str:  # SQL alias / CTE name for the source table
        return self.source_table.lower()


@dataclass
class Model:
    name: str
    columns: list[ColumnMap] = field(default_factory=list)
    attrs: dict[str, list[str]] = field(default_factory=lambda: {k: [] for k in MODEL_LEVEL})
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    contract: bool = False
    surrogate_key: str | None = None
    use_staging: bool = False

    def first(self, key: str, default: str = "") -> str:
        return self.attrs[key][0] if self.attrs[key] else default

    @property
    def sources(self) -> list[tuple[str, str]]:
        """Distinct (source_name, table) in first-seen order; first = driving table."""
        return list(OrderedDict.fromkeys((c.source_name, c.source_table.lower()) for c in self.columns))

    @property
    def aliases(self) -> list[str]:
        return [t for _, t in self.sources]

    @property
    def materialization(self) -> str:
        return self.first("MATERIALIZATION", "view").lower()

    @property
    def unique_key(self) -> list[str]:
        keys: list[str] = []
        for k in self.attrs["UNIQUE_KEY"]:
            keys += [p.strip().lower() for p in k.split(",") if p.strip()]
        return list(OrderedDict.fromkeys(keys))

    @property
    def effective_unique_key(self) -> list[str]:
        return [self.surrogate_key] if self.surrogate_key else self.unique_key

    @property
    def target_tables(self) -> list[str]:
        return list(OrderedDict.fromkeys(c.target_table for c in self.columns if c.target_table))

    @property
    def folder(self) -> str:
        return self.first("MODEL_FOLDER", "").strip("/ ").lower()

    @property
    def tags(self) -> list[str]:
        return list(OrderedDict.fromkeys(t.lower() for v in self.attrs["TAGS"] for t in split_list(v)))

    @property
    def join_predicates(self) -> list[str]:
        return [p.strip() for j in self.attrs["JOIN_CONDITION"] for p in re.split(r"[;\n]", j) if p.strip()]


# --------------------------------------------------------------------------- #
# Reading the mapping
# --------------------------------------------------------------------------- #
def read_rows(path: Path, sheet: str | None) -> list[dict]:
    if path.suffix.lower() in {".csv", ".txt"}:
        with path.open(newline="", encoding="utf-8-sig") as fh:
            raw = list(csv.reader(fh))
    else:
        import openpyxl  # pip install openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if sheet:
            ws = wb[sheet]
        elif DEFAULT_SHEET in wb.sheetnames:
            ws = wb[DEFAULT_SHEET]
        else:
            ws = wb.worksheets[0]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    if not raw:
        raise ValueError("Mapping file is empty")

    headers = [norm_header(h) for h in raw[0]]
    missing = [c for c in REQUIRED if c not in headers]
    if missing:
        raise ValueError(f"Mapping is missing required column(s): {missing}. Found: {headers}")

    rows = []
    for i, r in enumerate(raw[1:], start=2):
        rec = {h: clean(v) for h, v in zip(headers, r)}
        if not any(rec.values()):
            continue
        rec["_row"] = i
        rows.append(rec)
    return rows


def build_models(rows: list[dict], opts: Options) -> "OrderedDict[str, Model]":
    models: "OrderedDict[str, Model]" = OrderedDict()
    for r in rows:
        raw_name = r.get("MODEL_NAME", "")
        if not raw_name:
            log.error("Row %s: MODEL_NAME is blank - row skipped", r["_row"])
            continue
        name = re.sub(r"\.sql$", "", raw_name, flags=re.I).strip().lower()
        m = models.setdefault(name, Model(name=name))

        if not r.get("SOURCE_TABLE") or not r.get("SOURCE_COLUMN"):
            m.errors.append(f"Row {r['_row']}: SOURCE_TABLE/SOURCE_COLUMN is blank")
            continue
        m.columns.append(ColumnMap(
            row=r["_row"],
            source_name=(r.get("SOURCE_NAME") or opts.source_name).lower(),
            source_table=r["SOURCE_TABLE"],
            source_column=r["SOURCE_COLUMN"],
            target_table=r.get("TARGET_TABLE", ""),
            target_column=r.get("TARGET_COLUMN", "") or r["SOURCE_COLUMN"],
            target_datatype=r.get("TARGET_DATATYPE", ""),
            description=r.get("COLUMN_DESCRIPTION", ""),
            tests=r.get("COLUMN_TESTS", ""),
        ))
        for key in MODEL_LEVEL:
            v = r.get(key, "")
            if v and v not in m.attrs[key]:
                m.attrs[key].append(v)
    return models


# --------------------------------------------------------------------------- #
# Column helpers
# --------------------------------------------------------------------------- #
def is_expression(text: str) -> bool:
    return not IDENT_RE.match(text.strip().strip('"'))


def target_alias(c: ColumnMap) -> str:
    if is_expression(c.target_column):
        return c.source_column.lower()
    return c.target_column.strip().strip('"').lower()


def source_expression(c: ColumnMap, qualify: bool) -> str:
    col = c.source_column.lower()
    if is_expression(c.target_column):
        expr = c.target_column.strip()
        if qualify:
            expr = re.sub(rf"(?<![\w.]){re.escape(c.source_column)}(?![\w(])", f"{c.alias}.{col}", expr, flags=re.I)
        return expr
    return f"{c.alias}.{col}" if qualify else col


def column_sql(c: ColumnMap, qualify: bool) -> str:
    expr = source_expression(c, qualify)
    if c.target_datatype.strip():
        expr = f"cast({expr} as {c.target_datatype.strip().lower()})"
    return f"{expr} as {target_alias(c)}"


def referenced_columns(m: Model) -> "OrderedDict[tuple[str, str], OrderedDict[str, None]]":
    """Every source column the model needs: mapped columns + table.column refs in
    expressions, join conditions and incremental filters (feeds the staging layer)."""
    refs: "OrderedDict[tuple[str, str], OrderedDict[str, None]]" = OrderedDict()
    by_alias = {t: (s, t) for s, t in m.sources}
    for c in m.columns:
        refs.setdefault((c.source_name, c.alias), OrderedDict())[c.source_column.lower()] = None
    texts = [c.target_column for c in m.columns if is_expression(c.target_column)]
    texts += m.join_predicates + m.attrs["INCREMENTAL_FILTER"]
    for t in texts:
        t = re.sub(r"\{\{.*?\}\}", "", t)  # ignore jinja such as {{ this }}
        for alias, col in QUALIFIED_REF_RE.findall(t):
            if alias.lower() in by_alias:
                refs.setdefault(by_alias[alias.lower()], OrderedDict())[col.lower()] = None
    return refs


def parse_tests(spec: str) -> list:
    """not_null, unique, accepted_values(A|B), relationships(dim_x.col), any_pkg.test -> test entries."""
    tests: list = []
    for tok in split_list(spec):
        mt = re.match(r"^([\w.]+)\s*(?:\((.*)\))?$", tok.strip())
        if not mt:
            continue
        name, arg = mt.group(1).lower(), (mt.group(2) or "").strip()
        if name == "accepted_values" and arg:
            tests.append({name: {"values": [v.strip().strip("'\"") for v in re.split(r"[|;]", arg) if v.strip()]}})
        elif name == "relationships" and "." in arg:
            model, col = arg.rsplit(".", 1)
            tests.append({name: {"to": f"ref('{model.strip()}')", "field": col.strip().lower()}})
        else:
            tests.append(name)
    return tests


def merge_tests(*groups: list) -> list:
    out, seen = [], set()
    for g in groups:
        for t in g:
            key = t if isinstance(t, str) else next(iter(t))
            if key not in seen:
                seen.add(key)
                out.append(t)
    return out


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #
def aliases_in(text: str) -> set[str]:
    return {a.lower() for a, _ in QUALIFIED_REF_RE.findall(re.sub(r"\{\{.*?\}\}", "", text))}


def validate(m: Model, opts: Options) -> None:
    if not m.columns:
        m.errors.append("No usable column mappings")
        return

    mats = {x.lower() for x in m.attrs["MATERIALIZATION"]}
    if len(mats) > 1:
        m.errors.append(f"Conflicting MATERIALIZATION values: {sorted(mats)}")
    if not mats:
        m.warnings.append("MATERIALIZATION not set - defaulting to 'view'")
    elif m.materialization not in VALID_MATERIALIZATIONS:
        m.errors.append(f"Unsupported MATERIALIZATION '{m.materialization}'")
    for key in ("JOIN_TYPE", "MODEL_FOLDER", "SCHEMA", "MODEL_DESCRIPTION", "INCREMENTAL_FILTER", "USE_STAGING"):
        if len(m.attrs[key]) > 1:
            m.errors.append(f"Conflicting {key} values: {m.attrs[key]}")

    if not SNAKE_RE.match(m.name):
        m.warnings.append(f"Model name '{m.name}' is not snake_case")
    if len(m.target_tables) > 1:
        rows = ", ".join(f"row {c.row}={c.target_table}" for c in m.columns)
        m.warnings.append(f"Rows map to different TARGET_TABLEs {m.target_tables} under one MODEL_NAME ({rows}). "
                          "Check that MODEL_NAME is correct on every row.")

    seen: dict[str, int] = {}
    for c in m.columns:
        a = target_alias(c)
        if a in seen:
            m.errors.append(f"Duplicate target column '{a}' (rows {seen[a]} and {c.row})")
        seen[a] = c.row
        if not SNAKE_RE.match(a):
            m.warnings.append(f"Target column '{a}' (row {c.row}) is not snake_case")
    for k in m.unique_key:
        if k not in seen:
            m.errors.append(f"UNIQUE_KEY column '{k}' is not one of the model's target columns {list(seen)}")

    srcs = m.sources
    tables = [t for _, t in srcs]
    if len(set(tables)) != len(tables):
        m.errors.append(f"Same table name used from different SOURCE_NAMEs {srcs}; aliases would collide")
    if len(srcs) > 1:
        if not m.join_predicates:
            per = {t: [c.row for c in m.columns if c.alias == t] for t in tables}
            m.errors.append(f"Model has {len(srcs)} sources {tables} but no JOIN_CONDITION. "
                            f"Source rows: {per}. Add a JOIN_CONDITION or correct MODEL_NAME.")
        else:
            referenced = set().union(*(aliases_in(j) for j in m.join_predicates))
            for t in tables[1:]:
                if t not in referenced:
                    m.errors.append(f"Source '{t}' is not referenced by any JOIN_CONDITION {m.join_predicates}")
    elif m.join_predicates:
        m.warnings.append(f"JOIN_CONDITION supplied for single-source model - ignored: {m.join_predicates}")
    jt = m.first("JOIN_TYPE", "inner").lower()
    if jt not in VALID_JOIN_TYPES:
        m.errors.append(f"Unsupported JOIN_TYPE '{jt}'")

    if m.materialization == "incremental":
        if not m.unique_key:
            m.warnings.append("Incremental model without UNIQUE_KEY - dbt will append rather than merge")
        if not m.attrs["INCREMENTAL_FILTER"]:
            m.warnings.append("Incremental model without INCREMENTAL_FILTER - a TODO filter is generated")

    # Feature decisions -----------------------------------------------------
    flag = m.first("USE_STAGING").lower()
    if not flag:
        m.use_staging = opts.staging
    elif flag in YES or flag in NO:
        m.use_staging = flag in YES
    else:
        m.errors.append(f"USE_STAGING must be Y or N, got '{m.first('USE_STAGING')}'")
    if opts.surrogate_key and len(m.unique_key) > 1:
        m.surrogate_key = f"{m.name}_sk"
        if m.surrogate_key in seen:
            m.errors.append(f"Surrogate key name '{m.surrogate_key}' clashes with a mapped column")
    if opts.enforce_contracts:
        untyped = [target_alias(c) for c in m.columns if not c.target_datatype]
        if m.materialization == "ephemeral":
            m.warnings.append("Contracts are not supported on ephemeral models - contract skipped")
        elif untyped:
            m.warnings.append(f"Contract skipped: TARGET_DATATYPE missing for {untyped}")
        else:
            m.contract = True


# --------------------------------------------------------------------------- #
# SQL rendering
# --------------------------------------------------------------------------- #
def staging_name(source: str, table: str) -> str:
    return f"stg_{source}__{table}"


def render_staging(source: str, table: str, columns: list[str], opts: Options) -> str:
    cols = ",\n        ".join(columns)
    return (
        f"-- Staging model for {source}.{table}. Auto-generated by dbt_accelerator v{__version__} "
        f"from {opts.mapping_file}.\n"
        "-- One staging model per source table: light renaming only, no joins or business logic.\n\n"
        "with\n\n"
        "source as (\n\n"
        f"    select * from {{{{ source('{source}', '{table}') }}}}\n\n"
        "),\n\n"
        "renamed as (\n\n"
        "    select\n"
        f"        {cols}\n"
        "    from source\n\n"
        ")\n\n"
        "select * from renamed\n"
    )


def render_config(m: Model) -> str:
    parts = [f"materialized='{m.materialization}'"]
    keys = m.effective_unique_key
    if keys and m.materialization != "ephemeral":
        parts.append(f"unique_key='{keys[0]}'" if len(keys) == 1 else f"unique_key={keys!r}")
    if m.materialization == "incremental":
        parts += ["incremental_strategy='merge'", "on_schema_change='append_new_columns'"]
    if m.first("SCHEMA"):
        parts.append(f"schema='{m.first('SCHEMA').lower()}'")
    if m.tags:
        parts.append(f"tags={m.tags!r}")
    tgt = m.target_tables
    if len(tgt) == 1 and tgt[0].lower() != m.name:
        parts.append(f"alias='{tgt[0].lower()}'")
    if m.contract:
        parts.append("contract={'enforced': true}")
    return "{{\n    config(\n        " + ",\n        ".join(parts) + "\n    )\n}}"


def render_joins(m: Model) -> list[str]:
    tables = m.aliases
    jt = m.first("JOIN_TYPE", "inner").lower()
    preds = m.join_predicates
    lines, joined, pending = [], {tables[0]}, tables[1:]
    while pending:
        progressed = False
        for t in list(pending):
            ps = [p for p in preds if t in aliases_in(p) and aliases_in(p) <= joined | {t}]
            if ps:
                lines.append(f"    {jt} join {t}\n        on " + "\n        and ".join(ps))
                joined.add(t)
                pending.remove(t)
                progressed = True
        if not progressed:
            for t in pending:
                lines.append(f"    {jt} join {t}\n        on 1 = 1  -- TODO: JOIN_CONDITION missing for '{t}'")
            break
    return lines


def render_model(m: Model, opts: Options) -> str:
    qualify = len(m.sources) > 1
    select: list[str] = []
    if m.surrogate_key:
        key_exprs = []
        for k in m.unique_key:
            c = next(c for c in m.columns if target_alias(c) == k)
            key_exprs.append(f"'{source_expression(c, qualify)}'")
        select.append(f"{{{{ dbt_utils.generate_surrogate_key([{', '.join(key_exprs)}]) }}}} as {m.surrogate_key}")
    select += [column_sql(c, qualify) for c in m.columns]
    if opts.audit_columns:
        select += ["cast({{ dbt.current_timestamp() }} as timestamp) as dbt_loaded_at",
                   "cast('{{ invocation_id }}' as varchar(50)) as dbt_invocation_id"]

    imports = []
    for s, t in m.sources:
        rel = f"ref('{staging_name(s, t)}')" if m.use_staging else f"source('{s}', '{t}')"
        imports.append(f"{t} as (\n\n    select * from {{{{ {rel} }}}}\n\n)")

    body = ["    select", "        " + ",\n        ".join(select), f"    from {m.aliases[0]}"]
    body += render_joins(m) if len(m.sources) > 1 else []
    if m.materialization == "incremental":
        flt = m.first("INCREMENTAL_FILTER") or \
            "1 = 1  -- TODO: e.g. updated_at > (select max(updated_at) from {{ this }})"
        body += ["", "    {% if is_incremental() %}", f"    where {flt}", "    {% endif %}"]

    header = [
        render_config(m), "",
        f"-- Auto-generated by dbt_accelerator v{__version__} from {opts.mapping_file}.",
        "-- Do not edit by hand: change the mapping document and regenerate.", "",
        "with", "",
    ]
    return "\n".join(header) + "\n" + ",\n\n".join(imports) + ",\n\nfinal as (\n\n" + "\n".join(body) + \
        "\n\n)\n\nselect * from final\n"


# --------------------------------------------------------------------------- #
# Minimal YAML emitter (no PyYAML dependency, deterministic output)
# --------------------------------------------------------------------------- #
_PLAIN = re.compile(r"^[A-Za-z_][A-Za-z0-9_ .()/\-]*$")


def yq(v) -> str:
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    s = str(v)
    if _PLAIN.match(s) and s.lower() not in {"true", "false", "yes", "no", "null", "on", "off"}:
        return s
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"') + '"'


def ydump(obj, indent: int = 0) -> list[str]:
    pad, lines = "  " * indent, []
    if isinstance(obj, dict):
        for k, v in obj.items():
            if isinstance(v, (dict, list)) and v:
                lines.append(f"{pad}{k}:")
                lines += ydump(v, indent + 1)
            elif isinstance(v, (dict, list)):
                lines.append(f"{pad}{k}: {'{}' if isinstance(v, dict) else '[]'}")
            else:
                lines.append(f"{pad}{k}: {yq(v)}")
    else:
        for item in obj:
            if isinstance(item, dict) and item:
                sub = ydump(item, indent + 1)
                lines.append(f"{pad}- {sub[0].lstrip()}")
                lines += sub[1:]
            elif isinstance(item, list):
                lines.append(f"{pad}- [{', '.join(yq(x) for x in item)}]")
            else:
                lines.append(f"{pad}- {yq(item)}")
    return lines


def yaml_doc(obj) -> str:
    return "\n".join(ydump(obj)) + "\n"


def test_entry(t, opts: Options):
    """Wrap test args under `arguments:` (dbt >= 1.10) unless legacy syntax requested."""
    if isinstance(t, str) or opts.legacy_test_syntax:
        return t
    name, args = next(iter(t.items()))
    return {name: {"arguments": args}}


# --------------------------------------------------------------------------- #
# YAML rendering
# --------------------------------------------------------------------------- #
def render_sources_yml(source: str, tables: "OrderedDict[str, OrderedDict[str, None]]", opts: Options) -> str:
    src: dict = {"name": source}
    if opts.source_schema and source == opts.source_name:
        src["schema"] = opts.source_schema
    if opts.freshness_field:
        src["config"] = {"loaded_at_field": opts.freshness_field,
                         "freshness": {"warn_after": {"count": 24, "period": "hour"},
                                       "error_after": {"count": 48, "period": "hour"}}}
    src["tables"] = [{"name": t, "columns": [{"name": c} for c in cols]} for t, cols in tables.items()]
    return yaml_doc({"version": 2, "sources": [src]})


def render_staging_yml(source: str, tables: "OrderedDict[str, OrderedDict[str, None]]") -> str:
    models = [{"name": staging_name(source, t),
               "description": f"Staging model for {source}.{t} (1:1 with source, renamed only).",
               "columns": [{"name": c} for c in cols]} for t, cols in tables.items()]
    return yaml_doc({"version": 2, "models": models})


def render_models_yml(models: list[Model], opts: Options) -> str:
    tk = opts.tests_key
    docs = []
    for m in models:
        keys = m.effective_unique_key
        entry: dict = {"name": m.name,
                       "description": m.first("MODEL_DESCRIPTION") or
                       f"Target {', '.join(m.target_tables) or m.name}; built from {', '.join(m.aliases)}."}
        cfg: dict = {"meta": {"generated_by": f"dbt_accelerator v{__version__}", "mapping_file": opts.mapping_file}}
        if m.contract:
            cfg["contract"] = {"enforced": True}
        entry["config"] = cfg
        if len(m.unique_key) > 1 and not m.surrogate_key and opts.composite_test:
            entry[tk] = [test_entry({"dbt_utils.unique_combination_of_columns":
                                     {"combination_of_columns": m.unique_key}}, opts)]
        cols = []
        if m.surrogate_key:
            col = {"name": m.surrogate_key, "description": f"Surrogate key: hash of {', '.join(m.unique_key)}."}
            if m.contract:
                col["data_type"] = "varchar"
            col[tk] = ["not_null", "unique"]
            cols.append(col)
        for c in m.columns:
            name = target_alias(c)
            col = {"name": name}
            if c.target_datatype:
                col["data_type"] = c.target_datatype.lower()
            src = f"{c.source_name}.{c.alias}.{c.source_column.lower()}"
            col["description"] = c.description or \
                (f"Mapped from {src}" + (f" via {c.target_column.strip()}" if is_expression(c.target_column) else ""))
            auto = []
            if name in m.unique_key:
                auto = ["not_null"] + (["unique"] if len(m.unique_key) == 1 else [])
            tests = merge_tests(auto, parse_tests(c.tests))
            if tests:
                col[tk] = [test_entry(t, opts) for t in tests]
            cols.append(col)
        if opts.audit_columns:
            cols.append({"name": "dbt_loaded_at", "description": "Timestamp the row was written by dbt.",
                         **({"data_type": "timestamp"} if m.contract else {})})
            cols.append({"name": "dbt_invocation_id", "description": "dbt invocation_id of the run that wrote the row.",
                         **({"data_type": "varchar(50)"} if m.contract else {})})
        entry["columns"] = cols
        docs.append(entry)
    return yaml_doc({"version": 2, "models": docs})


def render_packages(needs_utils: bool) -> str | None:
    if not needs_utils:
        return None
    return yaml_doc({"packages": [{"package": "dbt-labs/dbt_utils", "version": [">=1.1.0", "<2.0.0"]}]})


def render_project_snippet(folders: list[str], opts: Options, has_staging: bool) -> str:
    tree: dict = {}
    if has_staging:
        tree["staging"] = {"+materialized": "view", "+tags": ["staging"]}
    for f in folders:
        node = tree
        for part in f.split("/"):
            node = node.setdefault(part, {})
        node.setdefault("+materialized", "table")
    head = "# Merge into dbt_project.yml - folder-level defaults (model-level config() still wins).\n"
    return head + yaml_doc({"models": {opts.project_name: tree}})


def render_sqlfluff(opts: Options) -> str:
    return f"""# SQL lint config matching the generated style. Requires: pip install sqlfluff sqlfluff-templater-dbt
[sqlfluff]
templater = dbt
dialect = {opts.dialect}
max_line_length = 120
exclude_rules = ST06

[sqlfluff:templater:dbt]
project_dir = .

[sqlfluff:indentation]
tab_space_size = 4

[sqlfluff:rules:capitalisation.keywords]
capitalisation_policy = lower

[sqlfluff:rules:capitalisation.identifiers]
extended_capitalisation_policy = lower

[sqlfluff:rules:capitalisation.functions]
extended_capitalisation_policy = lower

[sqlfluff:rules:aliasing.table]
aliasing = explicit

[sqlfluff:rules:aliasing.column]
aliasing = explicit
"""


# --------------------------------------------------------------------------- #
# Generation
# --------------------------------------------------------------------------- #
def generate(models: "OrderedDict[str, Model]", out: Path, opts: Options, force: bool) -> tuple[list, list]:
    written, skipped = [], []
    for m in models.values():
        validate(m, opts)
        for w in m.warnings:
            log.warning("[%s] %s", m.name, w)
        for e in m.errors:
            log.error("[%s] %s", m.name, e)
        (skipped if m.errors and not force else written).append(m)

    mdir = out / "models"
    # Staging layer + sources ------------------------------------------------
    # by_source: every source table any model reads (-> sources.yml)
    # by_stage : only tables read by models that use the staging layer (-> stg_ models)
    by_source: "OrderedDict[str, OrderedDict[str, OrderedDict[str, None]]]" = OrderedDict()
    by_stage: "OrderedDict[str, OrderedDict[str, OrderedDict[str, None]]]" = OrderedDict()
    for m in written:
        for (s, t), cols in referenced_columns(m).items():
            by_source.setdefault(s, OrderedDict()).setdefault(t, OrderedDict()).update(cols)
            if m.use_staging:
                by_stage.setdefault(s, OrderedDict()).setdefault(t, OrderedDict()).update(cols)
    for s, tables in by_source.items():
        sdir = mdir / "staging" / s
        sdir.mkdir(parents=True, exist_ok=True)
        (sdir / f"_{s}__sources.yml").write_text(render_sources_yml(s, tables, opts), encoding="utf-8")
    for s, tables in by_stage.items():
        sdir = mdir / "staging" / s
        (sdir / f"_{s}__models.yml").write_text(render_staging_yml(s, tables), encoding="utf-8")
        for t, cols in tables.items():
            (sdir / f"{staging_name(s, t)}.sql").write_text(render_staging(s, t, list(cols), opts), encoding="utf-8")
            log.info("Wrote staging/%s/%s.sql", s, staging_name(s, t))

    # Target models -----------------------------------------------------------
    by_folder: "OrderedDict[str, list[Model]]" = OrderedDict()
    for m in written:
        by_folder.setdefault(m.folder or opts.default_folder, []).append(m)
    for folder, ms in by_folder.items():
        fdir = mdir / folder
        fdir.mkdir(parents=True, exist_ok=True)
        for m in ms:
            (fdir / f"{m.name}.sql").write_text(render_model(m, opts), encoding="utf-8")
            log.info("Wrote %s/%s.sql (%s, %d cols, %s%s)", folder, m.name, m.materialization, len(m.columns),
                     "via staging" if m.use_staging else "direct from source", ", contract" if m.contract else "")
        leaf = folder.split("/")[-1]
        (fdir / f"_{leaf}__models.yml").write_text(render_models_yml(ms, opts), encoding="utf-8")

    # Project scaffolding ------------------------------------------------------
    needs_utils = any(m.surrogate_key or (len(m.unique_key) > 1 and opts.composite_test) for m in written)
    pk = render_packages(needs_utils)
    if pk:
        (out / "packages.yml").write_text(pk, encoding="utf-8")
    (out / "dbt_project_snippet.yml").write_text(render_project_snippet(list(by_folder), opts, bool(by_stage)), encoding="utf-8")
    (out / ".sqlfluff").write_text(render_sqlfluff(opts), encoding="utf-8")
    write_report(models, written, skipped, by_source, by_stage, out, opts)
    return written, skipped


def write_report(models, written, skipped, by_source, by_stage, out: Path, opts: Options) -> None:
    feats = [("Staging default", opts.staging), ("Model contracts", opts.enforce_contracts),
             ("Surrogate keys", opts.surrogate_key), ("Audit columns", opts.audit_columns),
             ("Source freshness", bool(opts.freshness_field))]
    rep = [f"# dbt accelerator v{__version__} - generation report", "",
           f"Mapping: `{opts.mapping_file}`  ", "Features: " + ", ".join(f"{n}={'on' if v else 'off'}" for n, v in feats),
           "", "## Models", "",
           "| Model | Status | Folder | Materialization | Unique key | Staging | Contract | Upstream | Issues |",
           "|---|---|---|---|---|---|---|---|---|"]
    for m in models.values():
        status = "SKIPPED" if m in skipped else ("WRITTEN (with errors)" if m.errors else "OK")
        ups = ", ".join(staging_name(s, t) if m.use_staging else f"source {s}.{t}" for s, t in m.sources)
        issues = "<br>".join([f"ERROR: {e}" for e in m.errors] + [f"WARN: {w}" for w in m.warnings]) or "-"
        rep.append(f"| {m.name} | {status} | {m.folder or opts.default_folder} | {m.materialization} | "
                   f"{', '.join(m.effective_unique_key) or '-'} | {'yes' if m.use_staging else 'no'} | "
                   f"{'yes' if m.contract else 'no'} | {ups} | {issues} |")
    rep += ["", "## Sources", "", "| Source | Table | Staging model | Columns used |", "|---|---|---|---|"]
    for s, tables in by_source.items():
        for t, cols in tables.items():
            stg = staging_name(s, t) if t in by_stage.get(s, {}) else "-"
            rep.append(f"| {s} | {t} | {stg} | {', '.join(cols)} |")
    (out / "generation_report.md").write_text("\n".join(rep) + "\n", encoding="utf-8")


def drift_check(models, target: Path, opts: Options, force: bool) -> int:
    """Generate into a temp dir and compare with the committed output (for CI)."""
    with tempfile.TemporaryDirectory() as tmp:
        tmpp = Path(tmp)
        logging.disable(logging.WARNING)
        generate(models, tmpp, opts, force)
        logging.disable(logging.NOTSET)
        diffs = []
        for f in sorted(p for p in tmpp.rglob("*") if p.is_file()):
            rel = f.relative_to(tmpp)
            other = target / rel
            if not other.exists():
                diffs.append(f"missing : {rel}")
            elif not filecmp.cmp(f, other, shallow=False):
                diffs.append(f"changed : {rel}")
        for d in diffs:
            log.error("DRIFT %s", d)
        if diffs:
            log.error("%d file(s) differ from the mapping - regenerate and commit.", len(diffs))
            return 1
        log.info("No drift: %s matches the mapping.", target)
        return 0


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Generate dbt models from a source-to-target mapping document.")
    ap.add_argument("mapping", type=Path, help="Mapping file (.xlsx or .csv)")
    ap.add_argument("-o", "--output", type=Path, default=Path("dbt_output"), help="Output directory")
    ap.add_argument("--sheet", help="Worksheet name (default: STM_MAPPING if present, else first sheet)")
    g = ap.add_argument_group("sources")
    g.add_argument("--source-name", default="raw", help="Default dbt source name (overridden by SOURCE_NAME column)")
    g.add_argument("--source-schema", help="Physical schema for the default source")
    g.add_argument("--freshness-field", help="loaded_at_field for source freshness checks (e.g. _loaded_at)")
    g = ap.add_argument_group("best practices")
    g.add_argument("--staging", action=argparse.BooleanOptionalAction, default=False,
                   help="Build a stg_<source>__<table> layer and ref() it (default: off - models select from "
                        "source() directly). The USE_STAGING column (Y/N) overrides this per model.")
    g.add_argument("--marts-dir", default="marts", help="Default folder for target models (MODEL_FOLDER overrides)")
    g.add_argument("--enforce-contracts", action="store_true", help="Add enforced model contracts (needs datatypes)")
    g.add_argument("--surrogate-key", action="store_true",
                   help="Composite keys -> dbt_utils.generate_surrogate_key column used as unique_key")
    g.add_argument("--audit-columns", action="store_true", help="Add dbt_loaded_at and dbt_invocation_id columns")
    g.add_argument("--no-composite-test", action="store_true", help="Skip unique_combination_of_columns test")
    g.add_argument("--legacy-test-syntax", action="store_true",
                   help="Emit 'tests:' with top-level args (dbt < 1.10) instead of 'data_tests:' + 'arguments:'")
    g = ap.add_argument_group("project")
    g.add_argument("--project-name", default="my_dbt_project", help="dbt project name for dbt_project_snippet.yml")
    g.add_argument("--dialect", default="snowflake", help="sqlfluff dialect (snowflake, bigquery, databricks, ...)")
    g = ap.add_argument_group("run mode")
    g.add_argument("--force", action="store_true", help="Also write models that fail validation (TODO markers)")
    g.add_argument("--check", action="store_true", help="CI mode: fail if OUTPUT differs from what the mapping generates")
    g.add_argument("--clean", action="store_true", help="Delete OUTPUT/models before generating (removes orphans)")
    ap.add_argument("-v", "--verbose", action="store_true")
    ap.add_argument("--version", action="version", version=__version__)
    a = ap.parse_args(argv)

    logging.basicConfig(level=logging.DEBUG if a.verbose else logging.INFO, format="%(levelname)-7s %(message)s")
    opts = Options(source_name=a.source_name.lower(), source_schema=a.source_schema, staging=a.staging,
                   default_folder=a.marts_dir.strip("/ ").lower(), enforce_contracts=a.enforce_contracts,
                   surrogate_key=a.surrogate_key, audit_columns=a.audit_columns,
                   composite_test=not a.no_composite_test, legacy_test_syntax=a.legacy_test_syntax,
                   freshness_field=a.freshness_field, project_name=a.project_name, dialect=a.dialect,
                   mapping_file=a.mapping.name)
    models = build_models(read_rows(a.mapping, a.sheet), opts)

    if a.check:
        return drift_check(models, a.output, opts, a.force)
    if a.clean and (a.output / "models").exists():
        shutil.rmtree(a.output / "models")
    a.output.mkdir(parents=True, exist_ok=True)
    written, skipped = generate(models, a.output, opts, a.force)
    log.info("Done: %d model(s) written, %d skipped. Report: %s",
             len(written), len(skipped), a.output / "generation_report.md")
    return 1 if skipped else 0


if __name__ == "__main__":
    sys.exit(main())
